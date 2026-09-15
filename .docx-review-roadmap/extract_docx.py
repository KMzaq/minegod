from pathlib import Path
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph


def blocks(parent):
    root = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    for child in root.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


path = Path(r"C:\Users\ADMIN\Downloads\Mythic_TRPG_기획서_v0.1.docx")
document = Document(path)
print(f"SECTIONS={len(document.sections)} TABLES={len(document.tables)} PARAGRAPHS={len(document.paragraphs)}")
for index, block in enumerate(blocks(document), 1):
    if isinstance(block, Paragraph):
        text = " ".join(block.text.split())
        if text:
            print(f"P{index:04d} [{block.style.name}] {text}")
    else:
        print(f"T{index:04d} rows={len(block.rows)} cols={len(block.columns)}")
        for row_index, row in enumerate(block.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            print(f"  R{row_index:03d} | " + " || ".join(cells))
